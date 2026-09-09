"""
Build an independent full-history cache for the dashboard.

Output:
    full_manager_history_data.js

This file is intentionally separate so Full Manager History does not recompute
splicing logic in the browser and does not affect other pages.
"""

from __future__ import annotations

import json
import os
from collections import defaultdict
from datetime import datetime

BASE_DIR = os.path.dirname(__file__)
POST_JS_PATH = os.path.join(BASE_DIR, "manager_returns_data.js")
PRE_JS_PATH = os.path.join(BASE_DIR, "pre_appointment_data.js")
OUTPUT_JS_PATH = os.path.join(BASE_DIR, "full_manager_history_data.js")
MAP_XLSX_PATH = os.path.join(BASE_DIR, "SAA_map_final.xlsx")


def load_js_var(file_path: str, var_name: str) -> dict:
    with open(file_path, "r", encoding="utf-8") as f:
        text = f.read()

    marker = f"var {var_name} = "
    start = text.find(marker)
    if start == -1:
        raise ValueError(f"Variable {var_name} not found in {file_path}")

    start += len(marker)
    end = text.rfind(";")
    if end == -1 or end <= start:
        raise ValueError(f"Could not parse JS payload from {file_path}")

    payload = text[start:end].strip()
    return json.loads(payload)


def month_key(dt: datetime) -> str:
    return f"{dt.year:04d}-{dt.month:02d}"


def build_post_monthly(post_data: dict) -> dict:
    out = {}
    for mgr in post_data.get("managers", []):
        series = post_data.get("data", {}).get(mgr, [])
        if len(series) < 2:
            continue

        manager_code = str(series[0].get("rbc_fa_code", "")).strip()
        if manager_code.endswith(".0"):
            manager_code = manager_code[:-2]
        if not manager_code:
            continue

        # Build month-end snapshot per stream id so return-index ratios
        # are calculated on consistent index bases across months.
        # Prefer managed_account_id; fallback to managed_account_key for legacy caches.
        month_last_by_stream = {}
        stream_month_count = defaultdict(int)
        for row in series:
            try:
                dt = datetime.fromisoformat(str(row.get("date", "")).replace("Z", ""))
            except Exception:
                continue

            row_code = str(row.get("rbc_fa_code", "")).strip()
            if row_code.endswith(".0"):
                row_code = row_code[:-2]
            if row_code != manager_code:
                continue

            ri = row.get("return_index")
            bi = row.get("relative_saa_index")
            # Keep post month-end points as soon as portfolio index is available.
            # Benchmark index can be missing and will be handled as nullable.
            if ri is None:
                continue

            key = row.get("managed_account_id")
            if key is None or str(key).strip() == "":
                key = row.get("managed_account_key")
            if key is None:
                continue
            key = str(key).strip()
            if not key:
                continue

            mkey = month_key(dt)
            if key not in month_last_by_stream:
                month_last_by_stream[key] = {}
            if mkey not in month_last_by_stream[key]:
                stream_month_count[key] += 1

            month_last_by_stream[key][mkey] = {
                "return_index": float(ri),
                "relative_saa_index": float(bi) if bi is not None else None,
            }

        if not month_last_by_stream:
            continue

        # Secondary check: use a single consistent return stream per RBC_FA_CODE.
        # Choose the stream with the broadest month coverage.
        primary_stream = max(
            month_last_by_stream.keys(),
            key=lambda k: (stream_month_count.get(k, 0), max(month_last_by_stream[k].keys()) if month_last_by_stream[k] else "")
        )

        month_last = month_last_by_stream[primary_stream]
        months = sorted(month_last)
        if len(months) < 2:
            continue

        monthly = []
        for i in range(1, len(months)):
            prev = month_last[months[i - 1]]
            curr = month_last[months[i]]
            if prev["return_index"] == 0:
                continue

            fund_ret = curr["return_index"] / prev["return_index"] - 1.0
            prev_bi = prev.get("relative_saa_index")
            curr_bi = curr.get("relative_saa_index")
            if prev_bi is not None and curr_bi is not None and prev_bi != 0:
                saa_ret = curr_bi / prev_bi - 1.0
            else:
                saa_ret = None
            monthly.append({
                "month": months[i],
                "fund_ret": fund_ret,
                "saa_ret": saa_ret,
                "source": "post",
            })

        if monthly:
            out[mgr] = monthly

    return out


def build_pre_monthly_by_code(pre_data: dict, allowed_codes: set[str] | None = None) -> dict:
    by_code_strategy = defaultdict(lambda: defaultdict(list))
    for row in pre_data.get("data", []):
        code = str(row.get("ID_CODE", "")).strip()
        strat = str(row.get("STRATEGY_NAME", "")).strip()
        if not code or not strat:
            continue
        if allowed_codes is not None and code not in allowed_codes:
            continue

        date_val = row.get("DATE")
        if not date_val:
            continue

        try:
            dt = datetime.fromisoformat(str(date_val).replace("Z", ""))
        except Exception:
            continue

        fund = row.get("FUND_1M")
        saa = row.get("SAA_1M")
        if fund is None:
            continue

        by_code_strategy[code][strat].append(
            {
                "month": month_key(dt),
                "fund_ret": float(fund) / 100.0,
                "saa_ret": float(saa) / 100.0 if saa is not None else 0.0,
                "source": "pre",
            }
        )

    out = {}
    for code, strat_map in by_code_strategy.items():
        if not strat_map:
            continue

        best_strat = max(strat_map, key=lambda s: len(strat_map[s]))
        rows = strat_map[best_strat]

        # De-dup by month.
        month_map = {}
        for r in sorted(rows, key=lambda x: x["month"]):
            month_map[r["month"]] = r

        out[code] = {
            "strategy": best_strat,
            "rows": [month_map[m] for m in sorted(month_map)],
        }

    return out


def splice_rows(pre_rows: list, post_rows: list) -> list:
    if not pre_rows and not post_rows:
        return []
    if not post_rows:
        return list(pre_rows)
    if not pre_rows:
        return list(post_rows)

    # Use pre history only before post monthly returns begin.
    # As soon as post has initial monthly-return observations, post becomes source of truth.
    first_post = post_rows[0]["month"]
    merged = [r for r in pre_rows if r["month"] < first_post] + list(post_rows)
    return merged


def load_saa_map_meta() -> dict:
    if not os.path.exists(MAP_XLSX_PATH):
        return {}

    try:
        import openpyxl
    except ImportError:
        return {}

    wb = openpyxl.load_workbook(MAP_XLSX_PATH, data_only=True)
    ws = wb[wb.sheetnames[0]]

    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    norm = {h.upper(): i for i, h in enumerate(headers)}

    def find_col(*candidates: str):
        for c in candidates:
            if c in norm:
                return norm[c]
        return None

    code_ix = find_col("RBC_FA_CODE", "RBC FA CODE", "RBC_FA", "RBC CODE")
    saa_ix = find_col("INDEX_COMPOSITION", "INDEX COMPOSITION")
    mgr_ix = find_col("MANAGER_NAME", "MANAGER NAME")
    if code_ix is None:
        return {}

    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        code = str(row[code_ix]).strip() if row[code_ix] is not None else ""
        if code.endswith(".0"):
            code = code[:-2]
        if code:
            out[code] = {
                "saa_name": str(row[saa_ix]).strip() if saa_ix is not None and row[saa_ix] is not None else "",
                "manager_name": str(row[mgr_ix]).strip() if mgr_ix is not None and row[mgr_ix] is not None else "",
            }
    return out


def build_full_history(post_data: dict, pre_data: dict) -> dict:
    post_monthly = build_post_monthly(post_data)
    saa_map_meta = load_saa_map_meta()
    mapped_codes = set(saa_map_meta.keys())

    # Match Post-Appointment page universe: only managers mapped by SAA_map_final.
    allowed_codes = set()
    for mgr in post_data.get("managers", []):
        mgr_rows = post_data.get("data", {}).get(mgr, [])
        if not mgr_rows:
            continue
        code = str(mgr_rows[0].get("rbc_fa_code", "")).strip()
        if code.endswith(".0"):
            code = code[:-2]
        if code and (not mapped_codes or code in mapped_codes):
            allowed_codes.add(code)

    # Keep pre-appointment history only for IDs that map to an existing
    # post-appointment RBC_FA_CODE universe.
    pre_by_code = build_pre_monthly_by_code(pre_data, allowed_codes=allowed_codes)

    out = {
        "managers": [],
        "meta": {},
        "data": {},
        "last_updated": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }

    for mgr in post_data.get("managers", []):
        mgr_rows = post_data.get("data", {}).get(mgr, [])
        if not mgr_rows:
            continue

        code = str(mgr_rows[0].get("rbc_fa_code", "")).strip()
        if code.endswith(".0"):
            code = code[:-2]
        if not code:
            continue
        if mapped_codes and code not in mapped_codes:
            continue

        post_rows = post_monthly.get(mgr, [])
        pre_info = pre_by_code.get(code)
        pre_rows = pre_info["rows"] if pre_info else []

        spliced = splice_rows(pre_rows, post_rows)
        if not spliced:
            continue

        fund_idx = 100.0
        saa_idx = 100.0
        series = []
        for r in spliced:
            fund_ret = float(r["fund_ret"])
            raw_saa = r.get("saa_ret")
            saa_ret = float(raw_saa) if raw_saa is not None else 0.0
            fund_idx *= 1.0 + fund_ret
            saa_idx *= 1.0 + saa_ret
            series.append(
                {
                    "date": r["month"],
                    "fund_ret": fund_ret,
                    "saa_ret": float(raw_saa) if raw_saa is not None else None,
                    "fund_idx": fund_idx,
                    "saa_idx": saa_idx,
                    "source": r.get("source", "spliced"),
                }
            )

        out["managers"].append(mgr)
        out["data"][mgr] = series
        manager_display_name = saa_map_meta.get(code, {}).get("manager_name") or mgr
        managed_account_name = str(mgr_rows[0].get("managed_account_name", "")).strip() if mgr_rows else ""
        if not managed_account_name:
            managed_account_name = mgr
        out["meta"][mgr] = {
            "displayName": manager_display_name,
            "code": code,
            "managedAccountName": managed_account_name,
            "preStrategy": pre_info["strategy"] if pre_info else "N/A",
            "saaName": saa_map_meta.get(code, {}).get("saa_name", "N/A") or "N/A",
        }

    out["managers"].sort(key=lambda m: out["meta"].get(m, {}).get("displayName", m))
    return out


def export_full_history_js(payload: dict, output_path: str) -> None:
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("// Auto-generated full-history cache - do not edit manually\n")
        f.write("// Last updated: " + payload.get("last_updated", "") + "\n")
        f.write("var CACHED_FULL_HISTORY_DATA = ")
        json.dump(payload, f)
        f.write(";\n")


if __name__ == "__main__":
    print("Loading post-appointment cache...")
    post = load_js_var(POST_JS_PATH, "CACHED_MANAGER_DATA")

    print("Loading pre-appointment cache...")
    pre = load_js_var(PRE_JS_PATH, "CACHED_PRE_APPOINTMENT_DATA")

    print("Building spliced full-history cache...")
    payload = build_full_history(post, pre)

    export_full_history_js(payload, OUTPUT_JS_PATH)
    print(f"Created: {OUTPUT_JS_PATH}")
    print(f"Managers: {len(payload['managers'])}")
